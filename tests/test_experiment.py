"""Experiments: conditions and repetitions from names, the experiment summary, outputs and combining."""

import math

import numpy as np
import pytest
import yaml
from PIL import Image
from scipy import stats

from microscount.bio.experiment import ExperimentDesign, combine_result_folders, summarise_experiment
from microscount.bio.pairing import (
    assign_conditions, condition_key, image_name, pair_files, scan_files, short_labels, split_field_number,
)
from microscount.cli import main
from microscount.synthetic import translocation_field


def _rgb(a, c):
    x = np.zeros(a.shape + (3,), np.uint8)
    x[..., c] = a
    return x


def _leica_field(folder, name, ratio, seed, shape=(192, 192)):
    """One field as a Leica LAS X export: blue _ch00, green _ch01, blank _ch02."""
    n, t, _ = translocation_field(shape=shape, ratio=ratio, seed=seed)
    Image.fromarray(_rgb(n, 2)).save(folder / f"{name}_ch00.tif")
    Image.fromarray(_rgb(t, 1)).save(folder / f"{name}_ch01.tif")
    Image.fromarray(np.ones(shape, np.uint8)).save(folder / f"{name}_ch02.tif")


def _experiment_tree(root, ratios=None):
    """Two repetitions ('…ASSAY-3', '…ASSAY-4'), conditions in the file names, one spelling difference."""
    ratios = ratios or {"vehicle": 1.0, "stimulus 30 min": 2.5}
    seed = 0
    for rep, spell in (("3", "stimulus 30 min"), ("4", "Stimulus 30min")):
        d = root / f"2026-01-15 ASSAY-{rep}"
        d.mkdir(parents=True)
        for cond, r in ratios.items():
            name = spell if cond == "stimulus 30 min" else cond
            for k in (1, 2):
                seed += 1
                _leica_field(d, f"15.01.2026 Assay-{rep}_{name}-{k}", r, seed)
    return root


# ---------------------------------------------------------------------- names


def test_names_to_conditions_and_field_numbers():
    assert image_name("15.01.2026 Assay-4_drug 10 uM +1h-1 _ch00.tif") == \
        "15.01.2026 Assay-4_drug 10 uM +1h-1"
    assert image_name("C1-img01.tif") == "img01" and image_name("img01_DAPI.tif") == "img01"
    assert split_field_number("stimulus 30 min-2") == ("stimulus 30 min", 2)
    assert split_field_number("vehicle (3)") == ("vehicle", 3)
    assert split_field_number("vehicle_pos04") == ("vehicle", 4)
    assert split_field_number("dose 10") == ("dose 10", None)  # a number after a space belongs to the condition
    assert condition_key("stimulus 30 min") == condition_key("Stimulus-30min")
    assert short_labels(["2026-01-15 ASSAY-3", "2026-01-15 ASSAY-4"]) == \
        {"2026-01-15 ASSAY-3": "3", "2026-01-15 ASSAY-4": "4"}
    assert short_labels(["rep1", "rep2", "rep10"]) == {"rep1": "1", "rep2": "2", "rep10": "10"}
    assert short_labels(["Monday", "Tuesday"]) == {"Monday": "Monday", "Tuesday": "Tuesday"}


def test_leica_repetitions_from_folders_and_conditions_from_names(tmp_path):
    root = _experiment_tree(tmp_path / "data")
    from microscount.core.imageio import list_images

    pairs, unpaired = pair_files(scan_files(list_images(root)))
    assert len(pairs) == 8 and all(p.confidence == "name" for p in pairs)
    assert all(u.note.startswith("blank") for u in unpaired) and len(unpaired) == 8
    assert {p.condition for p in pairs} == {"vehicle", "stimulus 30 min"}  # spellings merged
    assert {p.repetition for p in pairs} == {"3", "4"}
    ids = {p.field_id for p in pairs}
    assert "3 · vehicle #1" in ids and "4 · stimulus 30 min #2" in ids and len(ids) == 8


def test_conditions_from_folders(tmp_path):
    for rep in ("exp1", "exp2"):
        for cond, r in (("control", 1.0), ("treated", 2.0)):
            d = tmp_path / rep / cond
            d.mkdir(parents=True)
            for k in (1, 2):
                n, t, _ = translocation_field(shape=(128, 128), ratio=r, seed=k)
                Image.fromarray(n).save(d / f"img-{k}_dapi.png")
                Image.fromarray(t).save(d / f"img-{k}_gfp.png")
    from microscount.core.imageio import list_images

    pairs, _ = pair_files(scan_files(list_images(tmp_path)))
    assert {(p.repetition, p.condition) for p in pairs} == {("1", "control"), ("1", "treated"), ("2", "control"),
                                                           ("2", "treated")}
    assert assign_conditions(pairs, "name") == "name"  # forcing file names: every field is 'img'
    assert {p.condition for p in pairs} == {"img"}


# ---------------------------------------------------------------------- the summary


def _rows():
    """Three repetitions of control / A / B with known field medians and cells."""
    rng = np.random.default_rng(1)
    fields, cells = [], []
    medians = {"1": {"ctrl": [1.0, 1.2], "A": [2.0, 2.4], "B": [1.5, 1.7]},
               "2": {"ctrl": [0.9, 1.1], "A": [2.6, 2.2], "B": [1.4, 1.2]},
               "3": {"ctrl": [1.0, 1.0], "A": [2.1, 2.3], "B": [1.9, 1.3]}}
    for rep, conds in medians.items():
        for cond, vals in conds.items():
            for k, v in enumerate(vals, start=1):
                fid = f"{rep}·{cond}#{k}"
                ok = not (rep == "1" and cond == "A" and k == 2)  # one failed paper threshold
                fields.append({"field": fid, "condition": cond, "repetition": rep, "median_nc": v,
                               "paper_ratio": 0.5 * v + (5.0 if not ok else 0.0), "paper_ok": ok,
                               "nuclei_count": 50})
                for x in rng.normal(v, 0.3, 40):
                    cells.append({"field": fid, "condition": cond, "repetition": rep, "included": True,
                                  "nc_ratio": float(x), "cn_ratio": 1 / float(x), "nuc_corr": 10.0, "cyto_corr": 5.0})
    return fields, cells, medians


def test_experiment_summary_values():
    fields, cells, M = _rows()
    design = ExperimentDesign(control="ctrl", references={"B": "A"}, comparisons=[["ctrl", "A"], ["A", "B"]],
                              order=["ctrl", "A", "B"])
    exp = summarise_experiment(fields, cells, design)
    assert exp.order == ["ctrl", "A", "B"] and exp.repetitions == ["1", "2", "3"]
    by = {(r["repetition"], r["condition"]): r for r in exp.conditions}
    for rep in M:
        for cond, vals in M[rep].items():
            row = by[(rep, cond)]
            assert row["field_median_nc_mean"] == pytest.approx(np.mean(vals))
            assert row["field_median_nc_sd"] == pytest.approx(np.std(vals, ddof=1))
            ref = {"ctrl": "ctrl", "A": "ctrl", "B": "A"}[cond]
            assert row["fold_reference"] == ref
            assert row["fold_change"] == pytest.approx(np.mean(vals) / np.mean(M[rep][ref]))
            # responders: above the 95th percentile of the control's cells in the same repetition
            ctrl = [c["nc_ratio"] for c in cells if c["repetition"] == rep and c["condition"] == "ctrl"]
            mine = [c["nc_ratio"] for c in cells if c["repetition"] == rep and c["condition"] == cond]
            cut = np.percentile(ctrl, 95)
            assert row["responder_cutoff"] == pytest.approx(cut)
            assert row["responder_fraction"] == pytest.approx(np.mean(np.array(mine) > cut))
    # the failed paper field is left out of the paper mean
    r = by[("1", "A")]
    assert r["paper_fields_used"] == 1 and r["paper_fields_excluded"] == 1 and r["paper_ratio_mean"] == pytest.approx(1.0)
    # comparisons: Welch per repetition, paired across repetitions
    comp = {(c["repetition"], c["condition_a"], c["condition_b"]): c for c in exp.comparisons}
    for rep in M:
        c = comp[(rep, "A", "B")]
        assert c["welch_p"] == pytest.approx(stats.ttest_ind(M[rep]["A"], M[rep]["B"], equal_var=False).pvalue)
        assert c["difference_pct"] == pytest.approx(100 * (np.mean(M[rep]["B"]) / np.mean(M[rep]["A"]) - 1))
    allr = comp[("all", "A", "B")]
    a = [np.mean(M[r]["A"]) for r in M]
    b = [np.mean(M[r]["B"]) for r in M]
    assert allr["paired_p"] == pytest.approx(stats.ttest_rel(b, a).pvalue)
    assert allr["same_direction"] is True and allr["n_repetitions"] == 3
    # across repetitions
    o = {x["condition"]: x for x in exp.overall}
    assert o["A"]["mean"] == pytest.approx(np.mean(a)) and o["A"]["sd"] == pytest.approx(np.std(a, ddof=1))
    assert o["A"]["repetition 2"] == pytest.approx(np.mean(M["2"]["A"]))


def test_default_comparisons_two_repetitions_and_paper_only():
    fields, cells, M = _rows()
    two = [f for f in fields if f["repetition"] != "3"]
    exp = summarise_experiment(two, [c for c in cells if c["repetition"] != "3"], ExperimentDesign(control="ctrl"))
    pairs = {(c["condition_a"], c["condition_b"]) for c in exp.comparisons}
    assert pairs == {("ctrl", "A"), ("ctrl", "B")}  # every condition vs the control
    allr = next(c for c in exp.comparisons if c["repetition"] == "all" and c["condition_b"] == "A")
    assert math.isnan(allr["paired_p"]) and "3 repetitions" in allr["test"]
    assert any("2 repetitions" in n for n in exp.notes)
    # paper method only: field values are the paper ratios of the fields that passed the check
    paper_rows = [{k: v for k, v in f.items() if k != "median_nc"} for f in fields]
    exp = summarise_experiment(paper_rows, [], ExperimentDesign(control="ctrl"))
    assert exp.measure.startswith("paper")
    r = next(r for r in exp.conditions if r["repetition"] == "1" and r["condition"] == "A")
    assert "n_cells" not in r and r["paper_ratio_mean"] == pytest.approx(1.0)
    c = next(c for c in exp.comparisons if c["repetition"] == "1" and c["condition_b"] == "A")
    assert c["n_b"] == 1 and c["mean_b"] == pytest.approx(1.0)


def test_unknown_control_is_reported():
    fields, cells, _ = _rows()
    exp = summarise_experiment(fields, cells, ExperimentDesign(control="vehicle"))
    assert any("vehicle" in n for n in exp.notes) and not exp.comparisons
    assert all(r["responder_rule"] == "N/C above 1" for r in exp.conditions)


# ---------------------------------------------------------------------- runs, files, combining


def test_run_combine_and_command_line(tmp_path):
    root = _experiment_tree(tmp_path / "data")
    out = tmp_path / "joint"
    assert main(["bio", "translocation", str(root), "--control", "vehicle", "--out", str(out)]) == 0
    for f in ("summary.csv", "comparisons.csv", "per_condition.csv", "per_field.csv", "per_cell.csv", "results.xlsx",
              "summary.png", "cells.png", "settings.yaml", "CITATION.txt"):
        assert (out / f).exists(), f
    from openpyxl import load_workbook

    wb = load_workbook(out / "results.xlsx", read_only=True)
    assert wb.sheetnames[:6] == ["Read me", "Summary", "Comparisons", "Conditions", "Fields", "Cells"]
    doc = yaml.safe_load((out / "settings.yaml").read_text(encoding="utf-8"))
    assert doc["settings"]["control_condition"] == "vehicle"
    assert {i["repetition"] for i in doc["inputs"]} == {"3", "4"}
    joint = {(r["repetition"], r["condition"]): r for r in _csv(out / "per_condition.csv")}
    assert set(joint) == {("3", "vehicle"), ("3", "stimulus 30 min"), ("4", "vehicle"), ("4", "stimulus 30 min")}
    assert float(joint[("3", "stimulus 30 min")]["fold_change"]) > 1.5  # true ratios 2.5 vs 1.0
    # each repetition analysed on its own, then combined: the same values
    for rep in ("3", "4"):
        folder = next(root.glob(f"*-{rep}"))
        assert main(["bio", "translocation", str(folder), "--out", str(tmp_path / f"r{rep}")]) == 0
    assert main(["bio", "combine", str(tmp_path / "r3"), str(tmp_path / "r4"), "--control", "vehicle",
                 "--out", str(tmp_path / "comb")]) == 0
    comb = {(r["repetition"], r["condition"]): r for r in _csv(tmp_path / "comb" / "per_condition.csv")}
    assert set(comb) == set(joint)
    for k in joint:
        assert float(comb[k]["field_median_nc_mean"]) == pytest.approx(float(joint[k]["field_median_nc_mean"]))
        assert float(comb[k]["responder_fraction"]) == pytest.approx(float(joint[k]["responder_fraction"]))
    exp, _, _ = combine_result_folders([tmp_path / "r3", tmp_path / "r4"], ExperimentDesign(control="vehicle"))
    assert exp.repetitions == ["3", "4"]
    # a saved run repeats exactly
    assert main(["run", str(out / "settings.yaml"), "--out", str(tmp_path / "again")]) == 0
    again = {(r["repetition"], r["condition"]): r for r in _csv(tmp_path / "again" / "per_condition.csv")}
    assert again[("4", "vehicle")]["field_median_nc_mean"] == joint[("4", "vehicle")]["field_median_nc_mean"]


def _csv(path):
    import csv

    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))
