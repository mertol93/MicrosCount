"""Materials & Mechanics: SEM metadata, data bars, pore measures, samples and repetitions, summaries, runs."""

import heapq
import io
import itertools
import math
import time
from pathlib import Path

import numpy as np
import pytest
import tifffile
import yaml
from PIL import Image, ImageDraw
from scipy import ndimage as ndi
from scipy import stats
from skimage.morphology import local_minima

from microscount.cli import main
from microscount.core.experiment import GroupDesign, combine_repetitions
from microscount.core.imageio import load_image
from microscount.core.naming import group_key
from microscount.materials.experiment import (
    ImageSpec, assign_samples, combine_result_folders, file_pixel_size, label_images, scan_sem_files,
    summarise_samples,
)
from microscount.materials.porosity import (
    PorositySettings, analyse_sem, bwmorph_majority, detect_data_bar, detect_overlay_bar, grey_from_image,
    matlab_multithresh, measure_data_bar_scale, meyer_watershed, split_pores_watershed,
)
from microscount.synthetic import sem_image

DATA = Path(__file__).parent / "data"
FEI = ("[User]\r\nDate=10/01/2026\r\n[Beam]\r\nHV=5000\r\n[Scan]\r\nPixelWidth=4.5e-009\r\nPixelHeight=4.5e-009\r\n"
       "HorFieldsize=1.89e-006\r\n[Stage]\r\nWorkingDistance=0.0041\r\n[Image]\r\nResolutionX=420\r\nResolutionY=300\r\n"
       "[Detectors]\r\nNumber=1\r\nName=TLD\r\nMode=SE\r\n")
ZEISS = ("0\r\n0\r\nAP_PIXEL_SIZE\r\nPixel Size = 4.465 nm\r\nAP_ACTUALKV\r\nEHT = 5.00 kV\r\nAP_WD\r\nWD = 4.5 mm\r\n"
         "DP_DETECTOR_CHANNEL\r\nSignal A = InLens\r\nAP_MAG\r\nMag = 10.00 K X\r\n")


def _micrograph(h=300, w=420):
    g = load_image(DATA / "SEM1.jpg").data[0]
    return np.tile(g, (3, 3))[:h, :w].copy()


def _with_bar(img, rows=40, style="fei", scale=(60, 159)):
    """Append an SEM data bar: text-like strokes and a scale bar from x = scale[0] to scale[1]."""
    h, w = img.shape
    im = Image.fromarray(np.vstack([img, np.zeros((rows, w), np.uint8)]))
    d = ImageDraw.Draw(im)
    if style == "fei":
        for i in range(4):
            d.text((10 + 90 * i, h + 5), f"HV {i} kV", fill=255)
        d.rectangle([scale[0], h + rows - 14, scale[1], h + rows - 10], fill=255)
    else:  # two-tone bar: a white bar with a black panel, black scale bar on white
        d.rectangle([0, h, w - 1, h + rows - 1], fill=255)
        d.rectangle([3, h + 3, w // 2, h + rows - 4], fill=0)
        d.text((10, h + 8), "EHT = 5 kV", fill=255)
        d.rectangle([w // 2 + 20 + scale[0], h + rows - 14, w // 2 + 20 + scale[1], h + rows - 9], fill=0)
    return np.asarray(im)


# ---------------------------------------------------------------------- files and pixel sizes


def test_sem_metadata_pixel_size_and_data_bar(tmp_path):
    img = np.vstack([_micrograph(), np.zeros((40, 420), np.uint8)])  # 300 scanned rows + 40-row bar
    tifffile.imwrite(tmp_path / "fei.tif", img, extratags=[(34682, "s", 0, FEI, True)])
    tifffile.imwrite(tmp_path / "zeiss.tif", img, extratags=[(34118, "s", 0, ZEISS, True)])
    tifffile.imwrite(tmp_path / "ij.tif", img, imagej=True, resolution=(4.0, 4.0), metadata={"unit": "um"})
    tifffile.imwrite(tmp_path / "dpi.tif", img, resolution=(72, 72), resolutionunit="INCH")
    Image.fromarray(img).save(tmp_path / "plain.png")
    fei = load_image(tmp_path / "fei.tif")
    assert fei.pixel_size_um == pytest.approx(0.0045) and fei.pixel_size_source == "FEI / Thermo Fisher metadata"
    assert fei.metadata["data_bar_px"] == 40 and fei.metadata["voltage_kv"] == 5.0
    assert fei.metadata["working_distance_mm"] == pytest.approx(4.1) and fei.metadata["detector"] == "TLD (SE)"
    z = load_image(tmp_path / "zeiss.tif")
    assert z.pixel_size_um == pytest.approx(0.004465) and z.metadata["detector"] == "InLens"
    assert z.metadata["magnification"] == "10.00 K X"
    assert load_image(tmp_path / "ij.tif").pixel_size_um == pytest.approx(0.25)
    assert load_image(tmp_path / "dpi.tif").pixel_size_um is None  # the 72 dpi of image editors is not a pixel size
    # the same without reading the pixels (used when images are added)
    px, src, meta = file_pixel_size(tmp_path / "fei.tif")
    assert px == pytest.approx(0.0045) and src.startswith("FEI") and meta["data_bar_px"] == 40
    assert file_pixel_size(tmp_path / "plain.png") == (None, "", {})


def test_data_bar_detection_and_scale_bar():
    g = _micrograph()
    assert detect_data_bar(g) == 0
    for name in ("SEM1", "SEM2"):  # the MATLAB sample images have no bar
        assert detect_data_bar(grey_from_image(load_image(DATA / f"{name}.jpg"))[0]) == 0
    for style in ("fei", "two-tone"):
        a = _with_bar(g, 40, style)
        k = detect_data_bar(a)
        assert k == 40
        assert measure_data_bar_scale(a, k) == 100
        buf = io.BytesIO()
        Image.fromarray(a).save(buf, "JPEG", quality=85)
        j = np.asarray(Image.open(io.BytesIO(buf.getvalue())))
        assert detect_data_bar(j) == 40 and measure_data_bar_scale(j, 40) == 100
    sat = g.copy()
    sat[-60:, :250] = 255  # a bright region at the bottom is not a bar
    assert detect_data_bar(sat) == 0
    flat = np.vstack([g, np.zeros((40, g.shape[1]), np.uint8)])  # an empty strip below the sample: no text, no bar
    assert detect_data_bar(flat) == 0
    for img in [g, sat, flat] + [grey_from_image(load_image(DATA / f"{n}.jpg"))[0] for n in ("SEM1", "SEM2")]:
        assert detect_overlay_bar(img) == 0


def _jpeg(a, quality=75):
    buf = io.BytesIO()
    Image.fromarray(a).save(buf, "JPEG", quality=quality)
    return np.asarray(Image.open(io.BytesIO(buf.getvalue())))


def _xl30(img, bar=(360, 559)):
    """Philips XL30 style: two lines of white text written over the bottom of the micrograph, and on
    the first line a scale bar with end ticks from x = bar[0] to bar[1]."""
    h, w = img.shape
    txt = Image.new("L", (w // 2, 32), 0)
    d = ImageDraw.Draw(txt)
    d.text((4, 2), "Acc.V  Spot Magn   Det  WD", fill=255)
    d.text((290, 2), "1 um", fill=255)
    d.text((4, 16), "5.00 kV 3.0  20000x  SE  9.8  GTU", fill=255)
    text = np.asarray(txt.resize((w, 64), Image.NEAREST)) > 128  # twice the size: 2 px strokes
    out = img.copy()
    out[h - 64:][text] = 255
    out[h - 46:h - 44, bar[0]:bar[1] + 1] = 255
    out[h - 54:h - 36, bar[0]:bar[0] + 2] = 255
    out[h - 54:h - 36, bar[1] - 1:bar[1] + 1] = 255
    return out


def test_data_bars_in_figures_and_bars_written_over_the_image(tmp_path):
    g = _micrograph()
    # a figure: the image resized and recompressed blurs the bar's graphics
    for style, scale in (("fei", (60, 159)), ("two-tone", (60, 159))):
        a = _with_bar(g, 40, style, scale)
        small = np.asarray(Image.fromarray(a).resize((294, 238), Image.BILINEAR))  # 70%
        j = _jpeg(small)
        k = detect_data_bar(j)
        assert abs(k - 28) <= 2 and abs(measure_data_bar_scale(j, k) - 70) <= 2
    # a scale bar split by its label, |-- 5 um --|, is measured from end to end
    a = np.vstack([g, np.zeros((40, g.shape[1]), np.uint8)])
    im = Image.fromarray(a)
    d = ImageDraw.Draw(im)
    d.text((10, 305), "HV 5 kV  WD 4 mm", fill=255)
    d.rectangle([200, 326, 249, 329], fill=255)
    d.text((257, 321), "5 um", fill=255)
    d.rectangle([290, 326, 339, 329], fill=255)
    a = np.asarray(im)
    assert detect_data_bar(a) == 40 and measure_data_bar_scale(a, 40) == 140
    # a data bar written over the micrograph (XL30): found, measured and left out of the analysis
    x = _xl30(_micrograph(300, 640))
    assert detect_data_bar(x) == 0
    k = detect_overlay_bar(x)
    assert 300 - 244 <= k <= 72  # all of the text (from row 244), and not much more
    assert measure_data_bar_scale(x, k, overlay=True) == 200
    assert detect_overlay_bar(_jpeg(x, 85)) == k
    Image.fromarray(x).save(tmp_path / "xl30.png")
    res = analyse_sem(load_image(tmp_path / "xl30.png"), PorositySettings(pixel_size_um=0.005))
    assert res.summary["data_bar_px"] == k and res.summary["data_bar_source"] == "detected (text over the image)"


# ---------------------------------------------------------------------- watershed


def _meyer_one_by_one(image, markers):
    """Meyer's flooding pixel by pixel in MATLAB's order, to check the vectorised version."""
    h, w = image.shape
    out = markers.astype(int).copy()
    queued = out != 0
    nbrs = [(-1, -1), (0, -1), (1, -1), (-1, 0), (1, 0), (-1, 1), (0, 1), (1, 1)]  # column-major
    heap, count = [], itertools.count()

    def push(r, c, level):
        for dr, dc in nbrs:
            rr, cc = r + dr, c + dc
            if 0 <= rr < h and 0 <= cc < w and not queued[rr, cc]:
                queued[rr, cc] = True
                heapq.heappush(heap, (max(image[rr, cc], level), next(count), rr, cc))

    for c in range(w):
        for r in range(h):
            if markers[r, c]:
                push(r, c, -math.inf)
    while heap:
        level, _, r, c = heapq.heappop(heap)
        labels = {out[r + dr, c + dc] for dr, dc in nbrs if 0 <= r + dr < h and 0 <= c + dc < w} - {0}
        if len(labels) == 1:  # otherwise a line pixel: stays 0 and does not spread
            out[r, c] = labels.pop()
            push(r, c, level)
    return out


def test_watershed_matches_pixel_by_pixel_flooding():
    rng = np.random.default_rng(3)
    for trial in range(120):
        h, w = (int(v) for v in rng.integers(3, 30, 2))
        if trial % 3 == 0:  # plateaus and ties
            img = rng.integers(0, 4, (h, w)).astype(float)
        elif trial % 3 == 1:
            img = ndi.gaussian_filter(rng.random((h, w)), 2)
        else:  # a pore distance map, as in the porosity analysis
            pore = ndi.gaussian_filter(rng.random((h, w)), 1.5) > 0.5
            d = -ndi.distance_transform_cdt(pore, metric="taxicab").astype(float)
            img = ndi.median_filter(d, size=3, mode="constant")
        if trial % 2:
            markers = ndi.label(local_minima(img, connectivity=2, allow_borders=True), structure=np.ones((3, 3)))[0]
        else:  # any markers, touching ones too
            markers = np.where(rng.random((h, w)) < 0.05, rng.integers(1, 6, (h, w)), 0)
        assert np.array_equal(meyer_watershed(img, markers), _meyer_one_by_one(img, markers)), trial


def test_watershed_lines_and_large_flat_regions():
    img = np.array([[0, 1, 2, 1, 0]] * 3, float)
    ws = meyer_watershed(img, np.array([[1, 0, 0, 0, 2]] * 3))
    assert (ws[:, 2] == 0).all() and (ws[:, :2] == 1).all() and (ws[:, 3:] == 2).all()
    # an uncropped black data bar: one large pore with few minima (minutes with scikit-image 0.26)
    solid = ~(np.vstack([_micrograph(), np.zeros((40, 420), np.uint8)]) <= 90)
    t = time.perf_counter()
    pores = split_pores_watershed(solid)
    assert time.perf_counter() - t < 20
    assert pores[-1].any() and not (pores & solid).any()


# ---------------------------------------------------------------------- one image


def _bar_tiff(tmp_path, name="fei.tif"):
    img = _with_bar(_micrograph(), 40)
    tifffile.imwrite(tmp_path / name, img, extratags=[(34682, "s", 0, FEI, True)])
    return tmp_path / name


def test_pixel_size_order_data_bar_and_measures(tmp_path):
    p = _bar_tiff(tmp_path)
    img = load_image(p)
    s = PorositySettings(pixel_size_um=0.5)
    r = analyse_sem(img, s)  # the file's pixel size wins over the default
    S = r.summary
    assert S["pixel_size_um"] == pytest.approx(0.0045) and S["pixel_size_source"].startswith("FEI")
    assert S["data_bar_px"] == 40 and "metadata" in S["data_bar_source"] and S["image_height_px"] == 300
    assert analyse_sem(img, s, pixel_size_um=0.01).summary["pixel_size_source"] == "typed for this image"
    assert analyse_sem(img, PorositySettings(pixel_size_um=0.5, use_metadata_pixel_size=False)).summary[
        "pixel_size_um"] == 0.5
    # measures from the pores
    inc = [q for q in r.pores if q["included"]]
    d = np.array([q["equivalent_diameter_um"] for q in inc])
    a = np.array([q["area_px"] for q in inc], float)
    assert d == pytest.approx(2 * 0.0045 * np.sqrt(a / np.pi))
    assert S["mean_pore_diameter_um"] == pytest.approx(d.mean())
    assert S["median_pore_diameter_um"] == pytest.approx(np.median(d))
    assert S["d10_pore_diameter_um"] == pytest.approx(np.percentile(d, 10))
    assert S["d90_pore_diameter_um"] == pytest.approx(np.percentile(d, 90))
    assert S["area_weighted_mean_pore_diameter_um"] == pytest.approx((d * a).sum() / a.sum())
    assert S["mean_pore_radius_um"] == pytest.approx(d.mean() / 2)
    assert S["pore_density_per_um2"] == pytest.approx(len(r.pores) / (S["analysed_area_px"] * 0.0045 ** 2))
    assert all(0 < q["circularity"] <= 1 for q in inc)
    # the bar is left out exactly as cropping it by hand
    flat = np.asarray(img.data[0])[:300]
    Image.fromarray(flat).save(tmp_path / "cropped.png")
    ref = analyse_sem(load_image(tmp_path / "cropped.png"), PorositySettings(pixel_size_um=0.0045, data_bar="none"))
    assert S["porosity"] == ref.summary["porosity"] and S["n_pores"] == ref.summary["n_pores"]
    # detected without metadata, kept with a warning when switched off, manual
    Image.fromarray(np.asarray(img.data[0])).save(tmp_path / "bar.png")
    png = load_image(tmp_path / "bar.png")
    assert analyse_sem(png, PorositySettings()).summary["data_bar_source"] == "detected"
    off = analyse_sem(png, PorositySettings(data_bar="none"))
    assert off.summary["data_bar_px"] == 0 and any("data bar" in w for w in off.warnings)
    assert analyse_sem(png, PorositySettings(data_bar="manual", crop_bottom_px=50)).summary["image_height_px"] == 290


def test_unknown_pixel_size_fixed_threshold_and_old_settings(tmp_path):
    sem, _ = sem_image(seed=4)
    Image.fromarray(sem).save(tmp_path / "s.png")
    img = load_image(tmp_path / "s.png")
    r = analyse_sem(img, PorositySettings(n_thresholds=2))
    assert math.isnan(r.summary["mean_pore_diameter_um"]) and r.summary["mean_pore_diameter_px"] > 0
    assert any("pixel size unknown" in w for w in r.warnings)
    level = float(matlab_multithresh(sem, 2)[0])
    fixed = analyse_sem(img, PorositySettings(n_thresholds=2, pore_threshold_value=level))
    assert fixed.summary["porosity"] == r.summary["porosity"] and fixed.summary["threshold_source"] == "fixed value"
    # settings files of version 0.2 crop only what they asked for
    assert PorositySettings.from_dict({"crop_bottom_px": 0, "pixel_size_um": 0.459}).data_bar == "none"
    old = PorositySettings.from_dict({"crop_bottom_px": 60})
    assert (old.data_bar, old.crop_bottom_px) == ("manual", 60)
    assert PorositySettings().data_bar == "auto" and PorositySettings().pixel_size_um is None


# ---------------------------------------------------------------------- samples and repetitions


def _experiment(root, samples=(("CA", 0.10), ("CA-CNF1", 0.18)), reps=(1, 2), n=3):
    seed = 0
    for rep in reps:
        d = root / f"batch-{rep}"
        d.mkdir(parents=True)
        for name, por in samples:
            for k in range(1, n + 1):
                seed += 1
                img, _ = sem_image(porosity=por, seed=seed)
                Image.fromarray(img).save(d / f"Membrane_{name}_5kx_{k:02d}.png")
    return root


def test_samples_from_names_and_folders(tmp_path):
    from microscount.core.imageio import list_images

    root = _experiment(tmp_path / "a", n=2)
    specs = scan_sem_files(list_images(root))
    assert {(s.sample, s.repetition) for s in specs} == {("CA", "1"), ("CA-CNF1", "1"), ("CA", "2"), ("CA-CNF1", "2")}
    assert "2 · CA-CNF1 #2" in {s.image_id for s in specs}
    for rep in ("day1", "day2"):  # folders named after the samples
        for g in ("neat", "filled"):
            d = tmp_path / "b" / rep / g
            d.mkdir(parents=True)
            for k in (1, 2):
                Image.fromarray(sem_image(seed=k)[0]).save(d / f"img_{k}.tif")
    specs = scan_sem_files(list_images(tmp_path / "b"))
    assert {(s.sample, s.repetition) for s in specs} == {("neat", "1"), ("filled", "1"), ("neat", "2"), ("filled", "2")}
    assert assign_samples(specs, "name") == "name" and {s.sample for s in specs} == {"img"}
    label_images(specs)


def _png(path, a):
    Image.fromarray(a).save(path)
    return path


def _rows():
    """Three repetitions of three samples with known image values and pores."""
    M = {"1": {"ref": [10.0, 12.0], "A": [20.0, 22.0, 21.0], "B": [15.0, 14.0]},
         "2": {"ref": [11.0, 11.5], "A": [19.0, 23.0], "B": [16.0, 15.0]},
         "3": {"ref": [9.0, 10.0], "A": [18.0, 20.0], "B": [13.0, 14.0]}}
    images, pores = [], []
    for rep, groups in M.items():
        for g, vals in groups.items():
            for k, v in enumerate(vals, start=1):
                iid = f"{rep} · {g} #{k}"
                images.append({"image_id": iid, "image": f"{g}{k}.tif", "sample": g, "repetition": rep,
                               "porosity_percent": v, "mean_pore_diameter_um": v / 10, "median_pore_diameter_um": v / 20,
                               "pore_density_per_um2": 100 / v, "pixel_size_um": 0.01, "n_pores": 4})
                for j, dd in enumerate((v / 10, v / 8, v / 6, v / 4)):
                    pores.append({"image_id": iid, "sample": g, "repetition": rep, "pore": j + 1,
                                  "equivalent_diameter_um": dd, "area_px": 10 * (j + 1), "included": True})
    return images, pores, M


def test_summary_values():
    images, pores, M = _rows()
    design = GroupDesign(reference="ref", references={"B": "A"}, comparisons=[["ref", "A"], ["A", "B"]])
    ms = summarise_samples(images, pores, design, bins=10)
    assert ms.order == ["ref", "A", "B"] and ms.repetitions == ["1", "2", "3"] and ms.size_unit == "µm"
    per = {(r["repetition"], r["sample"]): r for r in ms.samples}
    for rep, groups in M.items():
        for g, vals in groups.items():
            row = per[(rep, g)]
            assert row["porosity_percent_mean"] == pytest.approx(np.mean(vals))
            assert row["porosity_percent_sd"] == pytest.approx(np.std(vals, ddof=1))
            ref = {"ref": "ref", "A": "ref", "B": "A"}[g]
            assert row["porosity_percent_change_pct"] == pytest.approx(100 * (np.mean(vals) / np.mean(M[rep][ref]) - 1))
            d = np.array([v / q for v in vals for q in (10, 8, 6, 4)])
            a = d ** 2  # pore areas follow from the equivalent diameters, whatever the pixel size
            assert row["pooled_median_diameter_µm"] == pytest.approx(np.median(d))
            assert row["pooled_area_weighted_mean_diameter_µm"] == pytest.approx((d * a).sum() / a.sum())
    summ = {r["sample"]: r for r in ms.summary}
    means = [np.mean(M[r]["A"]) for r in M]
    assert summ["A"]["porosity_percent_mean"] == pytest.approx(np.mean(means))
    assert summ["A"]["porosity_percent_sd"] == pytest.approx(np.std(means, ddof=1))
    assert summ["A"]["replicates"] == "repetitions" and summ["A"]["n_images"] == 7
    comp = {(c["key"], c["repetition"], c["sample_a"], c["sample_b"]): c for c in ms.comparisons}
    for rep in M:
        c = comp[("porosity_percent", rep, "A", "B")]
        assert c["welch_p"] == pytest.approx(stats.ttest_ind(M[rep]["A"], M[rep]["B"], equal_var=False).pvalue)
    allr = comp[("porosity_percent", "all", "A", "B")]
    a = [np.mean(M[r]["A"]) for r in M]
    b = [np.mean(M[r]["B"]) for r in M]
    assert allr["paired_p"] == pytest.approx(stats.ttest_rel(b, a).pvalue) and allr["same_direction"] is True
    assert ("mean_pore_diameter_um", "all", "ref", "A") in comp  # every measure is compared
    # distributions: fractions add up
    for g in ms.order:
        rows = [r for r in ms.distribution if r["sample"] == g and r["repetition"] == "all"]
        assert len(rows) == 10
        assert sum(r["pores_pct"] for r in rows) == pytest.approx(100)
        assert sum(r["area_pct"] for r in rows) == pytest.approx(100)
        assert rows[-1]["cumulative_pct"] == pytest.approx(100)


def test_one_repetition_and_unknown_pixel_sizes():
    images, pores, M = _rows()
    one = [r for r in images if r["repetition"] == "1"]
    ms = summarise_samples(one, [p for p in pores if p["repetition"] == "1"], GroupDesign(reference="ref"))
    a = {r["sample"]: r for r in ms.summary}["A"]
    assert a["replicates"] == "images" and a["porosity_percent_sd"] == pytest.approx(np.std(M["1"]["A"], ddof=1))
    assert {(c["sample_a"], c["sample_b"]) for c in ms.comparisons} == {("ref", "A"), ("ref", "B")}
    for r in one:
        r["pixel_size_um"] = ""
        r["mean_pore_diameter_px"] = r["mean_pore_diameter_um"] * 100
    ms = summarise_samples(one, [], GroupDesign(reference="ref"))
    assert ms.size_unit == "px" and ms.measures[1].key == "mean_pore_diameter_px"
    assert any("in pixels" in n for n in ms.notes)


# ---------------------------------------------------------------------- runs


def _csv(path):
    import csv

    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def test_run_combine_and_command_line(tmp_path):
    root = _experiment(tmp_path / "data")
    out = tmp_path / "joint"
    assert main(["materials", "porosity", str(root), "--pixel-size", "0.05", "--n-thresholds", "2", "--reference", "CA",
                 "--out", str(out)]) == 0
    for f in ("summary.csv", "comparisons.csv", "per_sample.csv", "per_image.csv", "per_pore.csv",
              "pore_size_distribution.csv", "results.xlsx", "summary.png", "distribution.png", "settings.yaml",
              "CITATION.txt"):
        assert (out / f).exists(), f
    from openpyxl import load_workbook

    assert load_workbook(out / "results.xlsx", read_only=True).sheetnames[:7] == [
        "Read me", "Summary", "Comparisons", "Samples", "Images", "Pores", "Distribution"]
    summ = {r["sample"]: r for r in _csv(out / "summary.csv")}
    assert float(summ["CA"]["porosity_percent_mean"]) == pytest.approx(10, abs=1)
    assert float(summ["CA-CNF1"]["porosity_percent_mean"]) == pytest.approx(18, abs=1.5)
    assert float(summ["CA-CNF1"]["porosity_percent_change_pct"]) > 50
    doc = yaml.safe_load((out / "settings.yaml").read_text(encoding="utf-8"))
    assert doc["settings"]["reference_sample"] == "CA" and {i["repetition"] for i in doc["inputs"]} == {"1", "2"}
    # each repetition alone, then combined: the same values
    for rep in (1, 2):
        assert main(["materials", "porosity", str(root / f"batch-{rep}"), "--pixel-size", "0.05", "--n-thresholds", "2",
                     "--out", str(tmp_path / "sep" / f"batch-{rep}")]) == 0
    assert main(["materials", "combine", str(tmp_path / "sep" / "batch-1"), str(tmp_path / "sep" / "batch-2"),
                 "--reference", "CA", "--out", str(tmp_path / "comb")]) == 0
    joint = {(r["repetition"], r["sample"]): r for r in _csv(out / "per_sample.csv")}
    comb = {(r["repetition"], r["sample"]): r for r in _csv(tmp_path / "comb" / "per_sample.csv")}
    assert set(comb) == set(joint)
    for k in joint:
        for col in ("porosity_percent_mean", "mean_pore_diameter_um_mean", "porosity_percent_change_pct"):
            assert float(comb[k][col]) == pytest.approx(float(joint[k][col]))
    ms, _, _ = combine_result_folders([tmp_path / "sep" / "batch-1", tmp_path / "sep" / "batch-2"])
    assert ms.repetitions == ["1", "2"]
    # a saved run repeats exactly
    assert main(["run", str(out / "settings.yaml"), "--out", str(tmp_path / "again")]) == 0
    again = {(r["repetition"], r["sample"]): r for r in _csv(tmp_path / "again" / "per_sample.csv")}
    assert again[("2", "CA")]["porosity_percent_mean"] == joint[("2", "CA")]["porosity_percent_mean"]
    # so does a combined result, whose combine.yaml also opens as the experiment design
    assert main(["run", str(tmp_path / "comb" / "combine.yaml"), "--out", str(tmp_path / "comb2")]) == 0
    assert _csv(tmp_path / "comb2" / "per_sample.csv") == _csv(tmp_path / "comb" / "per_sample.csv")
    from microscount.modules import load_settings

    assert load_settings(tmp_path / "comb" / "combine.yaml")[1].reference_sample == "CA"
    # a saved run given another repetition label: image ids follow
    assert main(["materials", "porosity", "--config", str(tmp_path / "sep" / "batch-1" / "settings.yaml"),
                 "--repetition", "7", "--out", str(tmp_path / "r7")]) == 0
    assert {r["image_id"].split(" #")[0] for r in _csv(tmp_path / "r7" / "per_image.csv")} == {"CA", "CA-CNF1"}
    assert {r["repetition"] for r in _csv(tmp_path / "r7" / "per_image.csv")} == {"7"}


def test_folders_at_any_depth(tmp_path):
    seed = 0
    for polymer in ("PES", "PSf"):
        for sample, por in (("neat", 0.10), ("filled", 0.18)):
            for rep in (1, 2):
                d = tmp_path / "SEM" / "2026" / polymer / sample / f"batch {rep}"
                d.mkdir(parents=True)
                for k in (1, 2):
                    seed += 1
                    Image.fromarray(sem_image(shape=(120, 160), porosity=por, seed=seed)[0]).save(d / f"img_{k}.png")
    out = tmp_path / "out"
    assert main(["materials", "porosity", str(tmp_path / "SEM"), "--pixel-size", "0.05", "--n-thresholds", "2",
                 "--out", str(out)]) == 0
    rows = _csv(out / "per_image.csv")
    assert len(rows) == 16
    assert {(r["sample"], r["repetition"]) for r in rows} == {
        (f"{p}/{s}", b) for p in ("PES", "PSf") for s in ("neat", "filled") for b in "12"}
    assert {r["n_repetitions"] for r in _csv(out / "summary.csv")} == {"2"}
    # only the folder itself
    assert main(["materials", "porosity", str(tmp_path / "SEM"), "--no-subfolders", "--out", str(tmp_path / "o2")]) == 2


def test_combining_result_folders_keeps_repetitions_apart():
    def rows(reps, samples, n=2):
        return [{"repetition": rep, "sample": g, "image_id": (f"{rep} · " if len(reps) > 1 else "") + f"{g} #{k}"}
                for rep in reps for g in samples for k in range(1, n + 1)]

    # analysed alone, one folder per repetition: the folders are the repetitions
    a, b = rows(["1"], ["CA", "CNF"]), rows(["1"], ["CA", "CNF"])
    assert combine_repetitions(["x/rep1/results", "x/rep2/results"], [a, b], [[], []], "sample", None, "image_id") == []
    assert {r["repetition"] for r in a} == {"1"} and {r["repetition"] for r in b} == {"2"}
    assert b[0]["image_id"] == "2 · CA #1"
    # two joint runs that both hold repetitions 1 and 2 of the same samples: four repetitions, with a note
    a, b = rows(["1", "2"], ["CA", "CNF"]), rows(["1", "2"], ["CA", "CNF"])
    pores = [dict(r) for r in b]
    notes = combine_repetitions(["run-A", "run-B"], [a, b], [[], pores], "sample", None, "image_id")
    assert {r["repetition"] for r in a + b} == {"A/1", "A/2", "B/1", "B/2"} and len(notes) == 2
    assert pores[-1]["image_id"] == b[-1]["image_id"] == "B/2 · CNF #2"
    # samples analysed separately: repetition 1 is one repetition with both samples
    a, b = rows(["1", "2"], ["CA"]), rows(["1", "2"], ["CNF"])
    assert combine_repetitions(["ca", "cnf"], [a, b], [[], []], "sample", None, "image_id") == []
    assert {(r["repetition"], r["sample"]) for r in a + b} == {(x, g) for x in "12" for g in ("CA", "CNF")}
    # labels given
    a, b = rows(["1", "2"], ["CA"]), rows(["1"], ["CA"])
    combine_repetitions(["p", "q"], [a, b], [[], []], "sample", ["day 1", "day 2"], "image_id")
    assert {r["repetition"] for r in a} == {"day 1"} and b[0]["image_id"] == "day 2 · CA #1"


def test_repetitions_and_samples_from_awkward_layouts(tmp_path):
    def specs(*paths):
        out = [ImageSpec(str(tmp_path / p)) for p in paths]
        assign_samples(out)
        label_images(out)
        return {(Path(s.path).relative_to(tmp_path).parent.as_posix(), s.sample, s.repetition) for s in out}

    names = [f"{g}_{k:02d}.tif" for g in ("Membrane_CA_5kx", "Membrane_CA-CNF1_5kx") for k in (1, 2)]
    # the folder of every repetition has the same name
    got = specs(*[f"day{d}/SEM/{n}" for d in (1, 2) for n in names])
    assert {(f, r) for f, _, r in got} == {("day1/SEM", "1"), ("day2/SEM", "2")}
    # sample folders inside same-named folders
    got = specs(*[f"day{d}/membranes/{g}/img_{k}.tif" for d in (1, 2) for g in ("neat", "filled") for k in (1, 2)])
    assert {(g, r) for _, g, r in got} == {("neat", "1"), ("filled", "1"), ("neat", "2"), ("filled", "2")}
    # a repetition whose folder holds one sample only
    got = specs(*[f"batch-{b}/{n}" for b in (1, 2) for n in names], "batch-3/Membrane_CA_5kx_01.tif",
                "batch-3/Membrane_CA_5kx_02.tif")
    assert {(g, r) for _, g, r in got} == {(g, b) for g in ("CA", "CA-CNF1") for b in "12"} | {("CA", "3")}
    # several studies in one tree: folders that share few samples are separate experiments, not
    # repetitions, and a sample name found in two of them takes its study's folder name
    got = specs(*[f"papers/{s}/Figure 2/{g}/{g}.tif" for s, gs in (("A 2023", ("neat", "x1", "x2")),
                                                                   ("B 2024", ("neat", "y1", "y2"))) for g in gs])
    assert {(g, r) for _, g, r in got} == {("A 2023/neat", "1"), ("x1", "1"), ("x2", "1"), ("B 2024/neat", "1"),
                                           ("y1", "1"), ("y2", "1")}
    # ... while repetitions hold mostly the same samples, whatever their folders are called, and
    # folders named like repetitions are repetitions even when their samples differ
    got = specs(*[f"{r}/{g}/{g}_1.tif" for r in ("first", "second") for g in ("neat", "x1", "x2")],
                "second/x3/x3_1.tif")
    assert {(g, r) for _, g, r in got} == {(g, r) for g in ("neat", "x1", "x2") for r in ("first", "second")} | {
        ("x3", "second")}
    got = specs(*[f"day{d}/{g}/{g}_1.tif" for d, gs in ((1, ("neat", "x1")), (2, ("neat", "x2"))) for g in gs])
    assert {(g, r) for _, g, r in got} == {("neat", "1"), ("x1", "1"), ("neat", "2"), ("x2", "2")}
    # spellings: separators and case do not count, scripts and signs do
    assert group_key("stimulus 30 min") == group_key("Stimulus-30min") and group_key("µm") == group_key("μm")
    assert len({group_key(g) for g in ("PES", "PES+", "PES-", "контроль", "мембрана", "PES-ı", "PES-ş")}) == 7


def test_result_folders_and_pixel_sizes_from_resolution_tags(tmp_path):
    from microscount.core.imageio import list_images
    from microscount.core.report import find_result_folders

    for d in ("run1", "run2", "microscount_combined_1"):
        (tmp_path / d).mkdir()
        (tmp_path / d / "per_image.csv").write_text("image\n", encoding="utf-8")
    (tmp_path / "run2" / "combine.yaml").write_text("{}", encoding="utf-8")
    assert find_result_folders(tmp_path, "per_image.csv") == [tmp_path / "run1"]
    Image.fromarray(np.zeros((8, 8), np.uint8)).save(tmp_path / "microscount_combined_1" / "summary.png")
    assert list_images(tmp_path) == []
    g, _ = sem_image(seed=1)
    tifffile.imwrite(tmp_path / "print.tif", g, resolution=(600, 600), resolutionunit="INCH")  # 42 µm: a print setting
    tifffile.imwrite(tmp_path / "fine.tif", g, resolution=(5000, 5000), resolutionunit="INCH")  # 5.08 µm
    assert file_pixel_size(tmp_path / "print.tif")[0] is None
    assert file_pixel_size(tmp_path / "fine.tif")[:2] == (pytest.approx(5.08), "TIFF resolution")
    fine = load_image(tmp_path / "fine.tif")
    assert analyse_sem(fine, PorositySettings(pixel_size_um=0.01)).summary["pixel_size_source"] == "default setting"
    assert analyse_sem(fine, PorositySettings()).summary["pixel_size_um"] == pytest.approx(5.08)


def test_bright_pores_unknown_pixel_size_and_twelve_bit(tmp_path):
    g, _ = sem_image(porosity=0.15, seed=6)
    inv = load_image(_png(tmp_path / "inv.png", 255 - g))
    dark = analyse_sem(load_image(_png(tmp_path / "g.png", g)), PorositySettings(n_thresholds=2))
    bright = analyse_sem(inv, PorositySettings(n_thresholds=2, pores_are_bright=True))
    assert bright.summary["porosity"] == dark.summary["porosity"]
    assert bright.summary["pore_threshold"] == 255 - dark.summary["pore_threshold"]  # on the image's own grey scale
    fixed = analyse_sem(inv, PorositySettings(pores_are_bright=True, pore_threshold_value=150))
    solid = bwmorph_majority(inv.data[0] < 150, 1)  # pores: grey 150 and above
    assert fixed.summary["pore_threshold"] == 150 and fixed.summary["porosity"] == pytest.approx(1 - solid.mean())
    S = dark.summary  # no pixel size: every size measure in pixels
    for k in ("sd", "d10", "d90", "area_weighted_mean", "largest"):
        assert math.isfinite(S[f"{k}_pore_diameter_px"]) and not math.isfinite(S[f"{k}_pore_diameter_um"])
    # 12-bit data in a 16-bit file: saturation at 4095, contrast on the 12-bit scale
    t = (g.astype(np.uint16) * 16).clip(0, 4095)
    t[:20] = 4095
    tifffile.imwrite(tmp_path / "t12.tif", t)
    r = analyse_sem(load_image(tmp_path / "t12.tif"), PorositySettings(n_thresholds=2))
    assert r.summary["saturated_fraction"] == pytest.approx((t >= 4095).mean())
    assert not any("low contrast" in w for w in r.warnings)


def test_checks_uneven_brightness_and_number_of_thresholds(tmp_path):
    from microscount.synthetic import sem_surface

    g, truth = sem_surface(seed=3, porosity=0.05, style="flat")
    assert abs(truth.mean() - 0.05) < 0.005
    even = analyse_sem(load_image(_png(tmp_path / "even.png", g)), PorositySettings())
    S = even.summary
    assert S["porosity_N4_percent"] == pytest.approx(S["porosity_percent"])
    assert S["porosity_N2_percent"] >= S["porosity_N4_percent"] >= S["porosity_N6_percent"]
    assert S["evened_overlap"] > 0.6 and not any("evened out" in w for w in even.warnings)
    # a darker half: the darkest class follows it rather than the pores
    dark = g.astype(float)
    dark[:, : g.shape[1] // 2] *= 0.55
    shaded = analyse_sem(load_image(_png(tmp_path / "shaded.png", np.round(dark).astype(np.uint8))), PorositySettings())
    assert shaded.summary["evened_overlap"] < 0.4 and any("evened out" in w for w in shaded.warnings)
    fixed = analyse_sem(load_image(tmp_path / "even.png"), PorositySettings(pore_threshold_value=100))
    assert "porosity_N4_percent" not in fixed.summary and math.isfinite(fixed.summary["evened_overlap"])


def test_check_against_traced_masks(tmp_path):
    from microscount.core.imageio import list_images
    from microscount.materials.check import check_image, mask_for, read_mask
    from microscount.synthetic import sem_surface

    g, truth = sem_surface(seed=5, porosity=0.05, style="flat")
    Image.fromarray(g).save(tmp_path / "flat.png")
    mask = truth.astype(np.uint8) * 255
    mask[:, :50] = 128  # a strip that was not traced
    Image.fromarray(mask).save(tmp_path / "flat_pores.png")
    assert [p.name for p in list_images(tmp_path)] == ["flat.png"]  # the mask is never analysed
    assert mask_for(tmp_path / "flat.png") == tmp_path / "flat_pores.png"
    pore, traced = read_mask(tmp_path / "flat_pores.png")
    assert not traced[:, :50].any() and traced[:, 50:].all() and (pore == (truth & traced)).all()
    row = check_image(tmp_path / "flat.png", PorositySettings(n_thresholds=2))  # flat pores: two thresholds find them
    assert row["overlap"] > 0.9 and row["pores_found"] > 0.95 and abs(row["size_ratio"] - 1) < 0.05
    assert row["porosity_traced_percent"] == pytest.approx(100 * truth[:, 50:].mean(), abs=1e-9)
    default = check_image(tmp_path / "flat.png")
    assert default["overlap"] < row["overlap"] and default["difference_pp"] < 0
    out = tmp_path / "scores.csv"
    assert main(["materials", "check", str(tmp_path), "--n-thresholds", "2", "--out", str(out)]) == 0
    assert out.read_text(encoding="utf-8-sig").splitlines()[0].startswith("image,mask,porosity_found_percent")
    Image.fromarray(mask[:-10]).save(tmp_path / "flat_pores.png")  # not the size of the image
    assert main(["materials", "check", str(tmp_path)]) == 2
