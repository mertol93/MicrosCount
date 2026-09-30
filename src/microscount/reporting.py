"""Batch runs and result files (CSV, Excel, overlays, figures, settings, citation)."""

from __future__ import annotations

import csv
import datetime as _dt
import math
import re
from pathlib import Path
from typing import Callable

import numpy as np

from . import __version__
from .citation import citation_text
from .imageio import load_image
from .porosity import PorosityResult, PorositySettings, analyse_sem, summarise_porosity
from .render import jet_map, label_colours, porosity_overlay, save_png, translocation_overlay
from .translocation import FieldResult, FieldSpec, TranslocationSettings, analyse_field, summarise_conditions

# chart tokens (reference palette, light surface)
INK = "#0b0b0b"
INK2 = "#52514e"
GRID = "#e6e5e1"
AXIS = "#bdbcb6"
SERIES = "#2a78d6"
SURFACE = "#ffffff"

Progress = Callable[[int, int, str], None]


def default_output_dir(inputs: list[str | Path]) -> Path:
    base = Path(inputs[0]).parent if inputs else Path.home()
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return base / f"microscount_results_{stamp}"


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")[:80] or "image"


# ---------------------------------------------------------------------- tables


def _columns(rows: list[dict], first: list[str] | None = None) -> list[str]:
    cols: list[str] = []
    for k in first or []:
        if any(k in r for r in rows) and k not in cols:
            cols.append(k)
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def _fmt(v):
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return f"{v:.10g}"
    if isinstance(v, (np.floating,)):
        return _fmt(float(v))
    if isinstance(v, (np.integer,)):
        return int(v)
    if v is None:
        return ""
    return v


def write_csv(path: Path, rows: list[dict], first: list[str] | None = None) -> Path:
    cols = _columns(rows, first)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads UTF-8
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _fmt(r.get(k)) for k in cols})
    return path


def write_xlsx(path: Path, sheets: dict[str, list[dict]], firsts: dict[str, list[str]] | None = None,
               notes: dict[str, str] | None = None) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if notes and name in notes:
            for line in notes[name].splitlines():
                ws.append([line])
            ws.column_dimensions["A"].width = 120
            continue
        cols = _columns(rows, (firsts or {}).get(name))
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_fmt(r.get(k)) if not isinstance(r.get(k), (list, dict)) else str(r.get(k)) for k in cols])
        ws.freeze_panes = "A2"
        for i, k in enumerate(cols, start=1):
            width = min(45, max(10, len(str(k)) + 2))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    wb.save(path)
    return path


def _settings_rows(settings) -> list[dict]:
    return [{"setting": k, "value": str(v)} for k, v in settings.to_dict().items()]


def write_settings_yaml(path: Path, module: str, settings, inputs: list[dict] | None = None) -> Path:
    import yaml

    doc = {
        "microscount_version": __version__,
        "module": module,
        "created": _dt.datetime.now().isoformat(timespec="seconds"),
        "settings": settings.to_dict(),
    }
    if inputs is not None:
        doc["inputs"] = inputs
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(doc, f, sort_keys=False, allow_unicode=True)
    return path


def load_settings_yaml(path: str | Path):
    """Return (module, settings object, inputs) from a saved settings file."""
    import yaml

    with open(path, encoding="utf-8") as f:
        doc = yaml.safe_load(f) or {}
    module = doc.get("module", "translocation")
    if module == "porosity":
        s = PorositySettings.from_dict(doc.get("settings", {}))
    else:
        s = TranslocationSettings.from_dict(doc.get("settings", {}))
    return module, s, doc.get("inputs") or []


# ---------------------------------------------------------------------- figures


def _style(ax):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(AXIS)
        ax.spines[side].set_linewidth(0.8)
    ax.tick_params(colors=INK2, labelsize=9, length=3, width=0.8)
    ax.yaxis.grid(True, color=GRID, linewidth=0.6, linestyle="-")
    ax.set_axisbelow(True)
    ax.set_facecolor(SURFACE)


def _figure(ncols: int, width_per: float = 4.2, height: float = 3.6):
    """A pyplot-free Figure (safe to draw from worker threads)."""
    from matplotlib.figure import Figure

    fig = Figure(figsize=(width_per * ncols, height), dpi=150)
    fig.patch.set_facecolor(SURFACE)
    axes = fig.subplots(1, ncols, squeeze=False)
    return fig, axes[0]


def _ref_line(ax, y: float, label: str):
    ax.axhline(y, color=INK2, linewidth=0.8, zorder=1)
    ax.annotate(label, xy=(1.0, y), xycoords=("axes fraction", "data"), xytext=(-2, 3),
                textcoords="offset points", ha="right", va="bottom", fontsize=8, color=INK2)


def _points_with_mean(ax, groups: list[str], values: list[np.ndarray], rng: np.random.Generator):
    for i, v in enumerate(values):
        v = v[np.isfinite(v)]
        if v.size == 0:
            continue
        x = i + rng.uniform(-0.12, 0.12, v.size)
        ax.scatter(x, v, s=30, color=SERIES, edgecolor=SURFACE, linewidth=1.2, zorder=3)
        m = v.mean()
        sd = v.std(ddof=1) if v.size > 1 else 0.0
        ax.errorbar(i + 0.28, m, yerr=sd, fmt="_", color=INK, markersize=12, capsize=3, linewidth=1.0, zorder=4)
    ax.set_xticks(range(len(groups)))
    ax.set_xticklabels(groups, rotation=20 if max((len(g) for g in groups), default=0) > 10 else 0,
                       ha="right" if max((len(g) for g in groups), default=0) > 10 else "center", color=INK)
    ax.set_xlim(-0.6, len(groups) - 0.4)


def translocation_figure(path: Path, results: list[FieldResult]) -> Path | None:
    if not results:
        return None
    conds = list(dict.fromkeys(r.condition or "(none)" for r in results))
    per_cell = any(r.method == "per_cell" for r in results)
    rng = np.random.default_rng(0)
    fig, axes = _figure(2 if per_cell else 1)
    k = 0
    if per_cell:
        ax = axes[0]
        _style(ax)
        data = []
        for c in conds:
            v = np.array([cell["ratio"] for r in results if (r.condition or "(none)") == c for cell in r.cells if cell["included"]])
            data.append(v[np.isfinite(v) & (v > 0)])
        for i, v in enumerate(data):
            if v.size == 0:
                continue
            show = v if v.size <= 1500 else rng.choice(v, 1500, replace=False)
            ax.scatter(i + rng.uniform(-0.25, 0.25, show.size), show, s=4, color=SERIES, alpha=0.25, linewidth=0, zorder=2)
        ax.boxplot([d if d.size else [np.nan] for d in data], positions=range(len(conds)), widths=0.55,
                   whis=(5, 95), showfliers=False, patch_artist=False,
                   medianprops={"color": INK, "linewidth": 1.6},
                   boxprops={"color": INK, "linewidth": 0.9}, whiskerprops={"color": INK, "linewidth": 0.9},
                   capprops={"color": INK, "linewidth": 0.9})
        for i, v in enumerate(data):
            if v.size:
                ax.annotate(f"{np.median(v):.2f}", xy=(i, np.percentile(v, 95)), xytext=(0, 4), textcoords="offset points",
                            ha="center", va="bottom", fontsize=8, color=INK2)
        ax.set_xticks(range(len(conds)))
        ax.set_xticklabels(conds, color=INK)
        _ref_line(ax, 1.0, "N = C")
        ax.set_ylabel("N/C ratio per cell", color=INK, fontsize=10)
        ax.set_title("Per-cell N/C ratio (box: IQR, whiskers: 5–95%)", color=INK, fontsize=10, loc="left")
        k = 1
    ax = axes[k]
    _style(ax)
    vals = [np.array([r.summary.get("paper_ratio", math.nan) for r in results if (r.condition or "(none)") == c]) for c in conds]
    _points_with_mean(ax, conds, vals, rng)
    _ref_line(ax, 1.0, "N = C")
    ax.set_ylabel("N/C ratio per field", color=INK, fontsize=10)
    ax.set_title("Whole-field ratio, paper method (mean ± SD)", color=INK, fontsize=10, loc="left")
    fig.tight_layout()
    fig.savefig(path, facecolor=SURFACE)
    return path


def pore_histogram_figure(path: Path, radii: np.ndarray, bins: int = 25, title: str = "") -> Path | None:
    radii = radii[np.isfinite(radii)]
    fig, axes = _figure(1, 5.0, 3.4)
    ax = axes[0]
    _style(ax)
    if radii.size:
        ax.hist(radii, bins=bins, color=SERIES, edgecolor=SURFACE, linewidth=1.0, zorder=2)
        m = radii.mean()
        ax.axvline(m, color=INK2, linewidth=0.8, zorder=3)
        ax.annotate(f"mean {m:.2f} µm", xy=(m, 1.0), xycoords=("data", "axes fraction"), xytext=(4, -12),
                    textcoords="offset points", fontsize=8, color=INK2)
    ax.set_xlabel("Equivalent pore radius (µm)", color=INK, fontsize=10)
    ax.set_ylabel("Number of pores", color=INK, fontsize=10)
    ax.set_title(title or "Pore size distribution", color=INK, fontsize=10, loc="left")
    fig.tight_layout()
    fig.savefig(path, facecolor=SURFACE)
    return path


def porosity_figure(path: Path, results: list[PorosityResult]) -> Path | None:
    if not results:
        return None
    conds = list(dict.fromkeys(r.condition or "(none)" for r in results))
    rng = np.random.default_rng(0)
    fig, axes = _figure(2)
    _style(axes[0])
    _points_with_mean(axes[0], conds, [np.array([100 * r.summary["porosity"] for r in results if (r.condition or "(none)") == c]) for c in conds], rng)
    axes[0].set_ylabel("Porosity (%)", color=INK, fontsize=10)
    axes[0].set_title("Porosity per image (mean ± SD)", color=INK, fontsize=10, loc="left")
    _style(axes[1])
    _points_with_mean(axes[1], conds, [np.array([r.summary["mean_pore_radius_um"] for r in results if (r.condition or "(none)") == c]) for c in conds], rng)
    axes[1].set_ylabel("Mean pore radius (µm)", color=INK, fontsize=10)
    axes[1].set_title("Mean pore radius per image (mean ± SD)", color=INK, fontsize=10, loc="left")
    fig.tight_layout()
    fig.savefig(path, facecolor=SURFACE)
    return path


# ---------------------------------------------------------------------- batch: translocation

FIELD_FIRST = [
    "field", "condition", "method", "nuclear_file", "target_file", "n_cells_analysed", "median_ratio",
    "geomean_ratio", "mean_ratio", "q1_ratio", "q3_ratio", "responder_fraction", "paper_ratio",
    "paper_nuclear_mean", "paper_cytoplasm_mean", "background", "nucleus_diameter_px",
]
CELL_FIRST = ["field", "condition", "cell", "included", "exclusion_reason", "ratio", "log2_ratio", "nuclear_mean", "cytoplasm_mean"]
COND_FIRST = ["condition", "n_fields", "n_cells", "pooled_median_ratio", "field_median_ratio_mean", "field_median_ratio_sd",
              "paper_ratio_mean", "paper_ratio_sd"]


def run_translocation(
    specs: list[FieldSpec],
    settings: TranslocationSettings,
    out_dir: str | Path,
    progress: Progress | None = None,
    cancel: Callable[[], bool] | None = None,
    save_overlays: bool = True,
) -> dict:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    ov_dir = out / "overlays"
    if save_overlays:
        ov_dir.mkdir(exist_ok=True)
    results: list[FieldResult] = []
    errors: list[dict] = []
    n = len(specs)
    for i, spec in enumerate(specs):
        if cancel and cancel():
            errors.append({"field": "", "error": "cancelled by user"})
            break
        fid = spec.field_id or f"field_{i + 1:03d}"
        if progress:
            progress(i, n, f"{fid}: {Path(spec.nuclear_path).name}")
        try:
            nuc = load_image(spec.nuclear_path)
            tgt = nuc if spec.target_path == spec.nuclear_path else load_image(spec.target_path)
            r = analyse_field(nuc, tgt, settings, fid, spec.condition, keep_layers=save_overlays)
            if save_overlays and r.layers is not None:
                save_png(ov_dir / f"{fid}_{_safe(Path(spec.nuclear_path).stem)}_overlay.png",
                         translocation_overlay(r.layers, base="target"))
            r.layers = None
            results.append(r)
        except Exception as exc:  # noqa: BLE001
            errors.append({"field": fid, "nuclear_file": spec.nuclear_path, "target_file": spec.target_path, "error": str(exc)})
    if progress:
        progress(n, n, "writing results")
    conditions = summarise_conditions(results)
    fields_rows = []
    for r in results:
        row = dict(r.summary)
        row["warnings"] = " | ".join(r.warnings)
        fields_rows.append(row)
    cell_rows = [c for r in results for c in r.cells]
    files = {
        "per_field": write_csv(out / "per_field.csv", fields_rows, FIELD_FIRST),
        "per_condition": write_csv(out / "per_condition.csv", conditions, COND_FIRST),
    }
    if cell_rows:
        files["per_cell"] = write_csv(out / "per_cell.csv", cell_rows, CELL_FIRST)
    if errors:
        files["errors"] = write_csv(out / "errors.csv", errors)
    inputs = [{"field": s.field_id, "condition": s.condition, "nuclear": s.nuclear_path, "target": s.target_path} for s in specs]
    files["settings"] = write_settings_yaml(out / "settings.yaml", "translocation", settings, inputs)
    cite = citation_text("translocation")
    (out / "CITATION.txt").write_text(cite, encoding="utf-8")
    files["citation"] = out / "CITATION.txt"
    sheets = {"Fields": fields_rows, "Conditions": conditions}
    if cell_rows:
        sheets["Cells"] = cell_rows
    if errors:
        sheets["Errors"] = errors
    sheets["Settings"] = _settings_rows(settings)
    sheets["How to cite"] = []
    x = write_xlsx(out / "results.xlsx", sheets, {"Fields": FIELD_FIRST, "Cells": CELL_FIRST, "Conditions": COND_FIRST},
                   notes={"How to cite": cite})
    if x:
        files["excel"] = x
    try:
        f = translocation_figure(out / "summary.png", results)
        if f:
            files["figure"] = f
    except Exception as exc:  # noqa: BLE001 - a plotting failure must not lose the results
        errors.append({"field": "", "error": f"summary figure failed: {exc}"})
    return {"results": results, "errors": errors, "conditions": conditions, "out_dir": out, "files": files}


# ---------------------------------------------------------------------- batch: porosity

POROSITY_FIRST = ["image", "condition", "porosity", "porosity_percent", "n_pores", "mean_pore_radius_um", "sd_pore_radius_um"]


def save_porosity_images(out: Path, stem: str, res: PorosityResult) -> dict:
    L = res.layers
    files = {}
    files["binary"] = out / f"{stem}_binary_segmentation.png"
    save_png(files["binary"], np.repeat((L["solid"].astype(np.uint8) * 255)[..., None], 3, axis=2))
    files["depth"] = out / f"{stem}_depth_map.png"
    save_png(files["depth"], jet_map(L["grey"]))
    files["pores"] = out / f"{stem}_pore_space_segmentation.png"
    save_png(files["pores"], label_colours(L["pore_labels"]))
    files["overlay"] = out / f"{stem}_overlay.png"
    save_png(files["overlay"], porosity_overlay(L["grey"], L["pore_labels"], L["valid"]))
    return files


def run_porosity(
    items: list[tuple[str, str]],  # (path, condition)
    settings: PorositySettings,
    out_dir: str | Path,
    progress: Progress | None = None,
    cancel: Callable[[], bool] | None = None,
    save_images: bool = True,
) -> dict:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    img_dir = out / "images"
    if save_images:
        img_dir.mkdir(exist_ok=True)
    results: list[PorosityResult] = []
    errors: list[dict] = []
    n = len(items)
    for i, (path, cond) in enumerate(items):
        if cancel and cancel():
            errors.append({"image": "", "error": "cancelled by user"})
            break
        if progress:
            progress(i, n, Path(path).name)
        try:
            img = load_image(path)
            r = analyse_sem(img, settings, cond, keep_layers=save_images)
            stem = f"{i + 1:03d}_{_safe(Path(path).stem)}"
            if save_images:
                save_porosity_images(img_dir, stem, r)
                radii = np.array([p["equivalent_radius_um"] for p in r.pores if p["included"]])
                pore_histogram_figure(img_dir / f"{stem}_pore_size_distribution.png", radii, settings.histogram_bins, Path(path).name)
            r.layers = None
            results.append(r)
        except Exception as exc:  # noqa: BLE001
            errors.append({"image": path, "error": str(exc)})
    if progress:
        progress(n, n, "writing results")
    summary_rows = []
    for r in results:
        row = dict(r.summary)
        row["warnings"] = " | ".join(r.warnings)
        summary_rows.append(row)
    pore_rows = []
    for r in results:
        for p in r.pores:
            pore_rows.append({"image": r.name, "condition": r.condition, **p})
    conditions = summarise_porosity(results)
    files = {
        "summary": write_csv(out / "porosity_summary.csv", summary_rows, POROSITY_FIRST),
        "pores": write_csv(out / "pores.csv", pore_rows),
        "per_condition": write_csv(out / "porosity_by_condition.csv", conditions),
    }
    if errors:
        files["errors"] = write_csv(out / "errors.csv", errors)
    inputs = [{"image": p, "condition": c} for p, c in items]
    files["settings"] = write_settings_yaml(out / "settings.yaml", "porosity", settings, inputs)
    cite = citation_text("porosity")
    (out / "CITATION.txt").write_text(cite, encoding="utf-8")
    files["citation"] = out / "CITATION.txt"
    sheets = {"Images": summary_rows, "Conditions": conditions, "Pores": pore_rows}
    if errors:
        sheets["Errors"] = errors
    sheets["Settings"] = _settings_rows(settings)
    sheets["How to cite"] = []
    x = write_xlsx(out / "results.xlsx", sheets, {"Images": POROSITY_FIRST}, notes={"How to cite": cite})
    if x:
        files["excel"] = x
    try:
        f = porosity_figure(out / "summary.png", results)
        if f:
            files["figure"] = f
    except Exception as exc:  # noqa: BLE001
        errors.append({"image": "", "error": f"summary figure failed: {exc}"})
    return {"results": results, "errors": errors, "conditions": conditions, "out_dir": out, "files": files}
