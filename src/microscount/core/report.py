"""Shared result writing: CSV, Excel, settings files and the figure style."""

from __future__ import annotations

import csv
import datetime as _dt
import math
import re
from pathlib import Path
from typing import Callable

import numpy as np

from .. import __version__

# chart tokens (reference palette, light surface)
INK = "#0b0b0b"
INK2 = "#52514e"
GRID = "#e6e5e1"
AXIS = "#bdbcb6"
SERIES = "#2a78d6"
SERIES_2 = "#eb6834"  # categorical slot 2 (validated adjacent pair with SERIES)
SURFACE = "#ffffff"

Progress = Callable[[int, int, str], None]

def default_output_dir(inputs: list[str | Path]) -> Path:
    base = Path(inputs[0]).parent if inputs else Path.home()
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return base / f"microscount_results_{stamp}"


def default_combined_dir(folders: list[str | Path]) -> Path:
    """A new folder for combined results, in the folder that holds all the result folders."""
    import os

    parents = [str(Path(f).resolve().parent) for f in folders]
    try:
        base = Path(os.path.commonpath(parents))
    except ValueError:  # different drives
        base = Path(parents[0])
    return base / ("microscount_combined_" + _dt.datetime.now().strftime("%Y%m%d_%H%M%S"))


def find_result_folders(root: str | Path, marker: str) -> list[Path]:
    """Result folders (holding ``marker``, e.g. ``per_image.csv``) at most two levels below ``root``,
    without combined results, which repeat the folders they were made from."""
    root = Path(root)
    found = {p.parent for pattern in (marker, f"*/{marker}", f"*/*/{marker}") for p in root.glob(pattern)}
    return sorted(f for f in found
                  if not (f / "combine.yaml").exists() and not f.name.lower().startswith("microscount_combined"))


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")[:80] or "image"


# ---------------------------------------------------------------------- tables


def _columns(rows: list[dict], first: list[str] | None = None) -> list[str]:
    cols: list[str] = []
    for k in first or []:
        if any(k in r for r in rows) and k not in cols:
            cols.append(k)
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def _fmt(v):
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return f"{v:.10g}"
    if isinstance(v, (np.floating,)):
        return _fmt(float(v))
    if isinstance(v, (np.integer,)):
        return int(v)
    if v is None:
        return ""
    return v


def write_csv(path: Path, rows: list[dict], first: list[str] | None = None) -> Path:
    cols = _columns(rows, first)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads UTF-8
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _fmt(r.get(k)) for k in cols})
    return path


def write_xlsx(path: Path, sheets: dict[str, list[dict]], firsts: dict[str, list[str]] | None = None,
               notes: dict[str, str] | None = None) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if notes and name in notes:
            for line in notes[name].splitlines():
                ws.append([line])
            ws.column_dimensions["A"].width = 120
            continue
        cols = _columns(rows, (firsts or {}).get(name))
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_fmt(r.get(k)) if not isinstance(r.get(k), (list, dict)) else str(r.get(k)) for k in cols])
        ws.freeze_panes = "A2"
        for i, k in enumerate(cols, start=1):
            width = min(45, max(10, len(str(k)) + 2))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    wb.save(path)
    return path


def _settings_rows(settings) -> list[dict]:
    return [{"setting": k, "value": str(v)} for k, v in settings.to_dict().items()]


def write_settings_yaml(path: Path, analysis: str, settings, inputs: list[dict] | None = None) -> Path:
    """Settings (and optionally the inputs) of one analysis; reopen with ``modules.load_settings``."""
    import yaml

    from ..modules import analysis as _analysis, module as _module

    a = _analysis(analysis)
    doc = {
        "microscount_version": __version__,
        "module": a.module,
        "module_title": _module(a.module).title,
        "analysis": a.key,
        "analysis_title": a.title,
        "created": _dt.datetime.now().isoformat(timespec="seconds"),
        "settings": settings.to_dict(),
    }
    if inputs is not None:
        doc["inputs"] = inputs
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(doc, f, sort_keys=False, allow_unicode=True)
    return path


def read_settings_yaml(path: str | Path) -> dict:
    """Raw content of a settings file written by ``write_settings_yaml``."""
    import yaml

    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


# ---------------------------------------------------------------------- figures


def _style(ax):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(AXIS)
        ax.spines[side].set_linewidth(0.8)
    ax.tick_params(colors=INK2, labelsize=9, length=3, width=0.8)
    ax.yaxis.grid(True, color=GRID, linewidth=0.6, linestyle="-")
    ax.set_axisbelow(True)
    ax.set_facecolor(SURFACE)


def _figure(ncols: int, width_per: float = 4.2, height: float = 3.6):
    """A pyplot-free Figure (safe to draw from worker threads)."""
    from matplotlib.figure import Figure

    fig = Figure(figsize=(width_per * ncols, height), dpi=150)
    fig.patch.set_facecolor(SURFACE)
    axes = fig.subplots(1, ncols, squeeze=False)
    return fig, axes[0]


def _ref_line(ax, y: float, label: str):
    ax.axhline(y, color=INK2, linewidth=0.8, zorder=1)
    ax.annotate(label, xy=(1.0, y), xycoords=("axes fraction", "data"), xytext=(-2, 3),
                textcoords="offset points", ha="right", va="bottom", fontsize=8, color=INK2)


def _points_with_mean(ax, groups: list[str], values: list[np.ndarray], rng: np.random.Generator):
    for i, v in enumerate(values):
        v = v[np.isfinite(v)]
        if v.size == 0:
            continue
        x = i + rng.uniform(-0.12, 0.12, v.size)
        ax.scatter(x, v, s=30, color=SERIES, edgecolor=SURFACE, linewidth=1.2, zorder=3)
        m = v.mean()
        sd = v.std(ddof=1) if v.size > 1 else 0.0
        ax.errorbar(i + 0.28, m, yerr=sd, fmt="_", color=INK, markersize=12, capsize=3, linewidth=1.0, zorder=4)
    ax.set_xticks(range(len(groups)))
    ax.set_xticklabels(groups, rotation=20 if max((len(g) for g in groups), default=0) > 10 else 0,
                       ha="right" if max((len(g) for g in groups), default=0) > 10 else "center", color=INK)
    ax.set_xlim(-0.6, len(groups) - 0.4)
